# -*- coding: utf-8 -*-
"""
Builds the reusable Chinese Test Generator question-bank template:
    templates/chinese-test-template.xlsx

The column layout defined here is the CONTRACT that js/core/excel-import.js
parses. If you change a header, change it in both places.

Run:  python3 tools/build_template.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "templates", "chinese-test-template.xlsx")

# ---- shared styles ---------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="C62828")      # required cols
OPT_FILL    = PatternFill("solid", fgColor="EF9A9A")      # optional cols
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT  = Font(bold=True, size=16, color="C62828")
NOTE_FONT   = Font(italic=True, color="555555")
WRAP        = Alignment(wrap_text=True, vertical="top")
CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN        = Side(style="thin", color="DDDDDD")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, headers, optional=()):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = OPT_FILL if h in optional else HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def write_rows(ws, rows, start=2):
    for r, row in enumerate(rows, start=start):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER


wb = Workbook()

# ===========================================================================
# 1. Instructions
# ===========================================================================
ws = wb.active
ws.title = "Instructions"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 2, "B": 110})

lines = [
    ("How to fill in this question bank", TITLE_FONT),
    ("", None),
    ("This one workbook is your whole question bank. Add, edit or delete rows, save it as .xlsx, "
     "then import it in the teacher app (index.html). You can re-import as often as you like.", NOTE_FONT),
    ("", None),
    ("GENERAL RULES", Font(bold=True, size=12)),
    ("• Dark-red headers are REQUIRED. Light-red headers are OPTIONAL — leave them blank if unused.", None),
    ("• Never rename or reorder the header row, and never rename the sheet tabs.", None),
    ("• Give every question a unique ID (QuestionID / SetID / PassageID). IDs keep answers correctly "
     "mapped even after the test is randomised, and let a future LMS track questions over time.", None),
    ("• Every question references a LessonID that exists on the 'Lessons' sheet.", None),
    ("• Type Pinyin WITH tone marks (nǐ hǎo, not ni3 hao3). The teacher decides at test time which "
     "content actually shows Pinyin; if a Pinyin cell is empty, that item simply shows none.", None),
    ("• Vietnamese is the meaning/translation language (MeaningVI column).", None),
    ("", None),
    ("THE SHEETS", Font(bold=True, size=12)),
    ("Lessons — define your lessons/units once. Every question points at one LessonID.", None),
    ("Pinyin — Section A. Show a Chinese word; student picks the correct Pinyin. Give the correct "
     "Pinyin plus three wrong ones (usually wrong tones).", None),
    ("Matching — Section B. Each row is ONE pair (Chinese ↔ Vietnamese meaning). Rows that share the "
     "same SetID form one matching activity. Aim for 4–8 pairs per SetID.", None),
    ("MultipleChoice — Section C. A prompt (a sentence with a gap, or line A of a dialogue) and up to "
     "four options. Put the correct option's LETTER (A/B/C/D) in the Correct column.", None),
    ("SentenceArrange — Section D. Write the sentence in the CORRECT order, splitting it into meaningful "
     "chunks with a forward slash:  我 / 每天 / 七点 / 起床 . The app shuffles the chunks for the student.", None),
    ("ReadingPassages + ReadingQuestions — Section E. Put each passage once on ReadingPassages with a "
     "PassageID. On ReadingQuestions, every question points back at its PassageID. A passage can have "
     "many questions; they always stay grouped with their passage.", None),
    ("", None),
    ("TIPS", Font(bold=True, size=12)),
    ("• Chunk pinyin / option pinyin must use the SAME number of chunks/options as the Chinese, in the "
     "same order, so they line up.", None),
    ("• The 'Correct' column accepts a letter (A) or the exact option text — a letter is safest.", None),
    ("• Add an Explanation to any question; it is shown to the student on the results page.", None),
    ("• Sample data is already filled in on every sheet. Replace it with your own.", None),
]
r = 1
for text, font in lines:
    cell = ws.cell(row=r, column=2, value=text)
    cell.alignment = WRAP
    if font:
        cell.font = font
    r += 1

# ===========================================================================
# 2. Lessons
# ===========================================================================
ws = wb.create_sheet("Lessons")
style_header(ws, ["LessonID", "LessonName", "Order"])
widths(ws, {"A": 14, "B": 34, "C": 10})
write_rows(ws, [
    ["L1", "Bài 1 – Chào hỏi (Greetings)", 1],
    ["L2", "Bài 2 – Số & Thời gian (Numbers & Time)", 2],
    ["L3", "Bài 3 – Địa điểm (Places)", 3],
    ["L4", "Bài 4 – Ăn uống (Food & Drink)", 4],
])

# ===========================================================================
# 3. Pinyin  (Section A)
# ===========================================================================
ws = wb.create_sheet("Pinyin")
headers = ["QuestionID", "LessonID", "Hanzi", "CorrectPinyin",
           "Distractor1", "Distractor2", "Distractor3", "Explanation"]
style_header(ws, headers, optional=["Explanation"])
widths(ws, {"A": 12, "B": 10, "C": 14, "D": 16, "E": 14, "F": 14, "G": 14, "H": 40})
write_rows(ws, [
    ["PY1", "L1", "中国", "zhōngguó", "zhòngguó", "zōngguó", "zhōnggǔo", "中 is first tone: zhōng."],
    ["PY2", "L1", "老师", "lǎoshī", "lāoshī", "láoshì", "lǎoshì", "老 third tone, 师 first tone."],
    ["PY3", "L1", "谢谢", "xièxie", "xièxiè", "xiéxie", "xièxue", "The second 谢 is neutral tone: xie."],
    ["PY4", "L3", "图书馆", "túshūguǎn", "tǔshūguǎn", "túshuguǎn", "túshūguàn", "馆 is third tone: guǎn."],
    ["PY5", "L2", "现在", "xiànzài", "xiǎnzài", "xiānzài", "xiànzǎi", "Both syllables are fourth tone."],
])

# ===========================================================================
# 4. Matching  (Section B)  — one pair per row, grouped by SetID
# ===========================================================================
ws = wb.create_sheet("Matching")
headers = ["SetID", "LessonID", "Chinese", "Pinyin", "MeaningVI", "Explanation"]
style_header(ws, headers, optional=["Pinyin", "Explanation"])
widths(ws, {"A": 10, "B": 10, "C": 16, "D": 18, "E": 24, "F": 34})
write_rows(ws, [
    ["MS1", "L1", "你好",  "nǐ hǎo",   "Xin chào", ""],
    ["MS1", "L1", "老师",  "lǎoshī",   "Giáo viên", ""],
    ["MS1", "L1", "学生",  "xuésheng", "Học sinh", ""],
    ["MS1", "L1", "谢谢",  "xièxie",   "Cảm ơn", ""],
    ["MS1", "L1", "再见",  "zàijiàn",  "Tạm biệt", ""],
    ["MS2", "L4", "米饭",  "mǐfàn",    "Cơm", ""],
    ["MS2", "L4", "水",    "shuǐ",     "Nước", ""],
    ["MS2", "L4", "茶",    "chá",      "Trà", ""],
    ["MS2", "L4", "苹果",  "píngguǒ",  "Táo", ""],
])

# ===========================================================================
# 5. MultipleChoice  (Section C)
# ===========================================================================
ws = wb.create_sheet("MultipleChoice")
headers = ["QuestionID", "LessonID", "Prompt", "PromptPinyin",
           "OptionA", "OptionAPinyin", "OptionB", "OptionBPinyin",
           "OptionC", "OptionCPinyin", "OptionD", "OptionDPinyin",
           "Correct", "Explanation"]
style_header(ws, headers,
             optional=["PromptPinyin", "OptionAPinyin", "OptionBPinyin",
                       "OptionCPinyin", "OptionDPinyin", "Explanation"])
widths(ws, {"A": 11, "B": 9, "C": 26, "D": 22,
            "E": 14, "F": 14, "G": 14, "H": 14, "I": 14, "J": 14, "K": 14, "L": 14,
            "M": 9, "N": 34})
write_rows(ws, [
    ["MC1", "L1", "A: 你叫什么名字？\nB: ______", "A: Nǐ jiào shénme míngzi?",
     "我叫李明。", "Wǒ jiào Lǐ Míng.", "我很好。", "Wǒ hěn hǎo.",
     "谢谢。", "Xièxie.", "再见。", "Zàijiàn.", "A",
     "The question asks for a name, so the reply gives a name."],
    ["MC2", "L1", "我 ______ 学生。", "Wǒ ______ xuésheng.",
     "是", "shì", "很", "hěn", "在", "zài", "有", "yǒu", "A",
     "是 links a noun to identity: 我是学生 (I am a student)."],
    ["MC3", "L2", "现在几点？ ______", "Xiànzài jǐ diǎn?",
     "三个", "sān gè", "三点", "sān diǎn", "三岁", "sān suì", "三本", "sān běn", "B",
     "点 is the measure word for clock time."],
    ["MC4", "L1", "A: 你好吗？\nB: ______", "A: Nǐ hǎo ma?",
     "我叫王芳。", "Wǒ jiào Wáng Fāng.", "我很好，谢谢。", "Wǒ hěn hǎo, xièxie.",
     "再见。", "Zàijiàn.", "不客气。", "Bú kèqi.", "B",
     "你好吗？ asks how you are; the natural reply is 我很好，谢谢。"],
])

# ===========================================================================
# 6. SentenceArrange  (Section D)
# ===========================================================================
ws = wb.create_sheet("SentenceArrange")
headers = ["QuestionID", "LessonID", "Chunks", "ChunkPinyin", "Meaning", "Explanation"]
style_header(ws, headers, optional=["ChunkPinyin", "Meaning", "Explanation"])
widths(ws, {"A": 11, "B": 9, "C": 30, "D": 34, "E": 34, "F": 30})
write_rows(ws, [
    ["SA1", "L2", "我 / 每天 / 七点 / 起床", "wǒ / měitiān / qī diǎn / qǐchuáng",
     "Tôi thức dậy lúc 7 giờ mỗi ngày.", "Time-of-day comes before the verb."],
    ["SA2", "L1", "她 / 是 / 我的 / 中文 / 老师", "tā / shì / wǒde / zhōngwén / lǎoshī",
     "Cô ấy là giáo viên tiếng Trung của tôi.", ""],
    ["SA3", "L4", "我们 / 明天 / 去 / 商店", "wǒmen / míngtiān / qù / shāngdiàn",
     "Ngày mai chúng tôi đi cửa hàng.", "Subject + time + verb + place."],
])

# ===========================================================================
# 7. ReadingPassages  (Section E, part 1)
# ===========================================================================
ws = wb.create_sheet("ReadingPassages")
headers = ["PassageID", "LessonID", "Title", "Passage", "PassagePinyin"]
style_header(ws, headers, optional=["Title", "PassagePinyin"])
widths(ws, {"A": 12, "B": 9, "C": 18, "D": 60, "E": 60})
write_rows(ws, [
    ["RP1", "L3", "我的一天",
     "我叫王芳。我是学生。我每天七点起床，八点去学校。我喜欢学习中文。"
     "下午我常常去图书馆看书。晚上我在家吃饭，然后做作业。",
     "Wǒ jiào Wáng Fāng. Wǒ shì xuésheng. Wǒ měitiān qī diǎn qǐchuáng, bā diǎn qù xuéxiào. "
     "Wǒ xǐhuan xuéxí zhōngwén. Xiàwǔ wǒ chángcháng qù túshūguǎn kàn shū. "
     "Wǎnshang wǒ zài jiā chīfàn, ránhòu zuò zuòyè."],
])

# ===========================================================================
# 8. ReadingQuestions  (Section E, part 2)
# ===========================================================================
ws = wb.create_sheet("ReadingQuestions")
headers = ["QuestionID", "PassageID", "Question", "QuestionPinyin",
           "OptionA", "OptionAPinyin", "OptionB", "OptionBPinyin",
           "OptionC", "OptionCPinyin", "OptionD", "OptionDPinyin",
           "Correct", "Explanation"]
style_header(ws, headers,
             optional=["QuestionPinyin", "OptionAPinyin", "OptionBPinyin",
                       "OptionCPinyin", "OptionDPinyin", "Explanation"])
widths(ws, {"A": 11, "B": 11, "C": 24, "D": 24,
            "E": 12, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12, "L": 12,
            "M": 9, "N": 32})
write_rows(ws, [
    ["RQ1", "RP1", "王芳几点起床？", "Wáng Fāng jǐ diǎn qǐchuáng?",
     "六点", "liù diǎn", "七点", "qī diǎn", "八点", "bā diǎn", "九点", "jiǔ diǎn", "B",
     "文中说“我每天七点起床”。"],
    ["RQ2", "RP1", "王芳是做什么的？", "Wáng Fāng shì zuò shénme de?",
     "老师", "lǎoshī", "学生", "xuésheng", "医生", "yīshēng", "司机", "sījī", "B",
     "“我是学生。”"],
    ["RQ3", "RP1", "下午王芳常常去哪儿？", "Xiàwǔ Wáng Fāng chángcháng qù nǎr?",
     "学校", "xuéxiào", "商店", "shāngdiàn", "图书馆", "túshūguǎn", "家", "jiā", "C",
     "“下午我常常去图书馆看书。”"],
])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("Wrote", os.path.abspath(OUT))
